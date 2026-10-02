import base64
import io
from copy import deepcopy
from difflib import SequenceMatcher
import json
import os
import tempfile
import time
import urllib.error
import urllib.request
from datetime import datetime
from pathlib import Path
from urllib.parse import urlparse

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Cm, Inches, Pt


CLOUD_RUN_SCREENSHOT_API = (
    "https://technical-seo-unlighthouse-api-231228645606."
    "asia-northeast1.run.app/screenshot"
)

PPT_FONT_FACE = "Meiryo UI"
DEFAULT_TEMPLATE_PATH = Path(__file__).resolve().parent / "templates" / "technical_seo_template.pptx"

SCREENSHOT_CACHE_TTL_SECONDS = 15 * 60
SCREENSHOT_CACHE_MAX_ITEMS = 40
_SCREENSHOT_CACHE = {}


def _post_json(api_url: str, payload: dict, timeout: int = 120):
    body = json.dumps(payload, ensure_ascii=False).encode("utf-8")

    request = urllib.request.Request(
        api_url,
        data=body,
        headers={"Content-Type": "application/json"},
        method="POST",
    )

    try:
        with urllib.request.urlopen(request, timeout=timeout) as response:
            return json.loads(response.read().decode("utf-8"))
    except urllib.error.HTTPError as e:
        error_body = e.read().decode("utf-8", errors="ignore")
        raise RuntimeError(
            f"Cloud Run HTTP Error {e.code}: {error_body}"
        ) from e


def _prune_screenshot_cache():
    now = time.time()

    expired = [
        url
        for url, item in _SCREENSHOT_CACHE.items()
        if now - item["created_at"] > SCREENSHOT_CACHE_TTL_SECONDS
    ]
    for url in expired:
        _SCREENSHOT_CACHE.pop(url, None)

    while len(_SCREENSHOT_CACHE) > SCREENSHOT_CACHE_MAX_ITEMS:
        oldest_url = next(iter(_SCREENSHOT_CACHE))
        _SCREENSHOT_CACHE.pop(oldest_url, None)


def _get_mobile_screenshot(url: str) -> bytes:
    _prune_screenshot_cache()

    cached = _SCREENSHOT_CACHE.get(url)
    if cached:
        # Refresh insertion order so repeated report generation reuses the
        # most recently requested screenshots without hitting Cloud Run again.
        _SCREENSHOT_CACHE.pop(url, None)
        _SCREENSHOT_CACHE[url] = cached
        return cached["bytes"]

    last_error = None

    # Cloud Run performs its own browser-level retries. Keep one additional
    # client retry for transient HTTP/container failures without multiplying
    # requests excessively for a 30-URL report.
    for delay_seconds in (0, 3):
        if delay_seconds:
            time.sleep(delay_seconds)

        try:
            data = _post_json(
                CLOUD_RUN_SCREENSHOT_API,
                {"url": url},
                timeout=120,
            )

            if not data.get("success"):
                raise RuntimeError(data.get("error") or "Screenshot failed")

            encoded = data.get("imageBase64")
            if not encoded:
                raise RuntimeError("Screenshot data was not returned")

            screenshot = base64.b64decode(encoded)
            if len(screenshot) < 3000:
                raise RuntimeError(
                    f"Screenshot data is unexpectedly small ({len(screenshot)} bytes)"
                )

            _SCREENSHOT_CACHE[url] = {
                "created_at": time.time(),
                "bytes": screenshot,
            }
            _prune_screenshot_cache()
            return screenshot
        except Exception as exc:
            last_error = exc

    raise RuntimeError(
        f"Screenshot failed after retries: {last_error}"
    ) from last_error


def _load_server_template_bytes() -> bytes:
    template_path = Path(os.getenv("PPT_TEMPLATE_PATH", str(DEFAULT_TEMPLATE_PATH)))

    if template_path.exists():
        return template_path.read_bytes()

    template_b64 = os.getenv("PPT_TEMPLATE_BASE64", "").strip()
    if template_b64:
        return base64.b64decode(template_b64)

    try:
        from default_ppt_template import DEFAULT_PPT_TEMPLATE_B64

        return base64.b64decode(DEFAULT_PPT_TEMPLATE_B64)
    except Exception:
        pass

    raise RuntimeError(
        "PowerPointテンプレートが見つかりません。"
        " テンプレートをアップロードするか、PPT_TEMPLATE_PATH / PPT_TEMPLATE_BASE64を設定してください。"
    )


def _set_cell_text(cell, text, font_size=6, bold=False, color="000000", align=None):
    cell.text = ""
    paragraph = cell.text_frame.paragraphs[0]
    if align:
        paragraph.alignment = align
    run = paragraph.add_run()
    run.text = str(text or "")
    run.font.name = PPT_FONT_FACE
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def _status_color(status: str):
    if status == "OK":
        return "003CFF"
    if status in {"NG", "△"}:
        return "D00000"
    return "555555"


def _replace_title_and_summary(slide, url: str, summary: str, show_summary: bool = True):
    parsed = urlparse(url)
    path = parsed.path or "/"
    date_label = f"{datetime.now().month}/{datetime.now().day}"
    title = f"{path}（{date_label}時点チェック）"

    summary = (summary or "").strip()
    if not summary:
        summary = "Technical SEOチェック結果を確認してください。"

    for shape in slide.shapes:
        if not getattr(shape, "has_text_frame", False):
            continue
        if "まとめ" not in shape.text and "URL" not in shape.text:
            continue

        text_frame = shape.text_frame
        text_frame.clear()

        p1 = text_frame.paragraphs[0]
        r1 = p1.add_run()
        r1.text = title
        r1.font.name = PPT_FONT_FACE
        r1.font.size = Pt(14)
        r1.font.bold = True
        r1.font.color.rgb = RGBColor.from_string("222222")

        if show_summary and summary:
            p2 = text_frame.add_paragraph()
            r2 = p2.add_run()
            r2.text = summary
            r2.font.name = PPT_FONT_FACE
            r2.font.size = Pt(14)
            r2.font.bold = True
            r2.font.color.rgb = RGBColor.from_string("0432FF")
        return


def _fill_table(slide, checks):
    table_shape = next((shape for shape in slide.shapes if getattr(shape, "has_table", False)), None)
    if table_shape is None:
        raise RuntimeError("Template table was not found")

    table = table_shape.table

    _set_cell_text(table.cell(0, 0), "No", 8, True, align=PP_ALIGN.CENTER)
    _set_cell_text(table.cell(0, 1), "チェック項目", 8, True, align=PP_ALIGN.CENTER)
    _set_cell_text(table.cell(0, 2), "", 8, True, align=PP_ALIGN.CENTER)
    _set_cell_text(table.cell(0, 3), "判定", 8, True, align=PP_ALIGN.CENTER)
    _set_cell_text(table.cell(0, 4), "結果", 8, True, align=PP_ALIGN.CENTER)

    max_rows = min(len(checks), len(table.rows) - 1)

    for index in range(max_rows):
        row = checks[index]
        ppt_row = index + 1

        result_text = str(row.get("Result") or "")
        action_text = str(row.get("Action") or "")
        if action_text and action_text != "対応不要":
            result_text += f"\nコメント: {action_text}"

        status = str(row.get("Status") or "")

        _set_cell_text(table.cell(ppt_row, 0), row.get("No"), 7, False, align=PP_ALIGN.CENTER)
        _set_cell_text(table.cell(ppt_row, 1), row.get("Item"), 7, True)
        _set_cell_text(table.cell(ppt_row, 2), row.get("Meaning"), 5, False)
        _set_cell_text(
            table.cell(ppt_row, 3),
            status,
            8,
            True,
            color=_status_color(status),
            align=PP_ALIGN.CENTER,
        )
        _set_cell_text(table.cell(ppt_row, 4), result_text, 7, False)

    for ppt_row in range(max_rows + 1, len(table.rows)):
        for col in range(len(table.columns)):
            table.cell(ppt_row, col).text = ""


def _add_screenshot_fixed(slide, screenshot_bytes: bytes):
    # Fixed placement based on the approved Technical SEO report layout.
    shot_left = Inches(0.55)
    shot_top = Inches(1.50)
    shot_width = Cm(6.5)
    shot_height = Inches(5.64)

    suffix = ".png" if screenshot_bytes.startswith(b"\x89PNG") else ".jpg"

    with tempfile.NamedTemporaryFile(
        suffix=suffix,
        delete=False,
    ) as image_file:
        image_file.write(screenshot_bytes)
        image_path = image_file.name

    try:
        slide.shapes.add_picture(
            image_path,
            shot_left,
            shot_top,
            width=shot_width,
            height=shot_height,
        )
    finally:
        try:
            os.remove(image_path)
        except FileNotFoundError:
            pass


def _style_proof_run(
    run,
    *,
    font_size=8,
    bold=False,
    color="000000",
):
    run.font.name = PPT_FONT_FACE
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def _add_original_with_diff(paragraph, original: str, suggestion: str):
    prefix = paragraph.add_run()
    prefix.text = "原文: "
    _style_proof_run(prefix)

    matcher = SequenceMatcher(
        None,
        original or "",
        suggestion or "",
    )

    for tag, i1, i2, _j1, _j2 in matcher.get_opcodes():
        segment = (original or "")[i1:i2]
        if not segment:
            continue

        run = paragraph.add_run()
        run.text = segment
        _style_proof_run(
            run,
            color="FF0000" if tag != "equal" else "000000",
        )


def _add_proofreading_box(
    slide,
    *,
    enabled: bool,
    proofreading,
    error: str = "",
):
    items = list(proofreading or [])

    # PPTには実際の指摘がある場合だけ誤字脱字欄を表示する。
    # 問題なし／GeminiエラーはWeb画面だけで表示する。
    if not enabled or error or not items:
        return

    # Coordinates and typography follow the user's approved sample.
    box = slide.shapes.add_textbox(
        Inches(3.4271325459),
        Inches(6.2312106299),
        Inches(9.4095374016),
        Inches(0.9087893701),
    )
    text_frame = box.text_frame
    text_frame.clear()
    text_frame.word_wrap = True
    text_frame.margin_left = Inches(0.10)
    text_frame.margin_right = Inches(0.10)
    text_frame.margin_top = Inches(0.05)
    text_frame.margin_bottom = Inches(0.05)

    note = text_frame.paragraphs[0]
    note.alignment = PP_ALIGN.LEFT

    run = note.add_run()
    run.text = "※誤字脱字"
    _style_proof_run(
        run,
        bold=True,
        color="FF0000",
    )

    run = note.add_run()
    run.text = (
        "：SEO判定には含めません。"
        "明確な誤字・脱字・変換ミスだけを確認します。"
    )
    _style_proof_run(run)

    max_items = 2

    for index, item in enumerate(items[:max_items], start=1):
        original = str(item.get("original") or "")
        suggestion = str(item.get("suggestion") or "")

        original_p = text_frame.add_paragraph()
        original_p.alignment = PP_ALIGN.LEFT

        number_run = original_p.add_run()
        number_run.text = f"{index}. "
        _style_proof_run(number_run)

        _add_original_with_diff(
            original_p,
            original,
            suggestion,
        )

        suggestion_p = text_frame.add_paragraph()
        suggestion_p.alignment = PP_ALIGN.LEFT
        suggestion_run = suggestion_p.add_run()
        suggestion_run.text = f"   修正案: {suggestion}"
        _style_proof_run(suggestion_run)

    if len(items) > max_items:
        more_p = text_frame.add_paragraph()
        more_p.alignment = PP_ALIGN.LEFT
        more_run = more_p.add_run()
        more_run.text = (
            f"ほか{len(items) - max_items}件は"
            "画面上の誤字脱字チェック結果を確認してください。"
        )
        _style_proof_run(
            more_run,
            font_size=7,
            color="555555",
        )


def _duplicate_template_slide(presentation, source_slide):
    new_slide = presentation.slides.add_slide(source_slide.slide_layout)

    for shape in list(new_slide.shapes):
        element = shape.element
        element.getparent().remove(element)

    for shape in source_slide.shapes:
        cloned = deepcopy(shape.element)
        new_slide.shapes._spTree.insert_element_before(cloned, "p:extLst")

    return new_slide


def build_multi_ppt_report_from_template(reports: list[dict], template_bytes: bytes):
    if not reports:
        raise RuntimeError("PowerPointに出力するレポートがありません。")

    presentation = Presentation(io.BytesIO(template_bytes))
    template_slide = presentation.slides[0]

    while len(presentation.slides) < len(reports):
        _duplicate_template_slide(presentation, template_slide)

    warnings = []

    for index, report in enumerate(reports):
        slide = presentation.slides[index]
        url = str(report.get("url") or "")
        checks = report.get("checks") or []
        summary = str(report.get("summary") or "")
        summary_enabled = bool(report.get("summary_enabled", True))
        proofreading = report.get("proofreading") or []
        proofreading_enabled = bool(
            report.get("proofreading_enabled", False)
        )
        proofreading_error = str(
            report.get("proofreading_error") or ""
        )

        _replace_title_and_summary(
            slide,
            url,
            summary,
            show_summary=summary_enabled,
        )
        _fill_table(slide, checks)
        _add_proofreading_box(
            slide,
            enabled=proofreading_enabled,
            proofreading=proofreading,
            error=proofreading_error,
        )

        try:
            screenshot = _get_mobile_screenshot(url)
            if len(screenshot) < 1000:
                raise RuntimeError(
                    f"Screenshot data is unexpectedly small ({len(screenshot)} bytes)"
                )
            _add_screenshot_fixed(slide, screenshot)
        except Exception as exc:
            warnings.append(f"{url}: モバイルスクリーンショット取得失敗 ({exc})")

    output = io.BytesIO()
    presentation.save(output)
    output.seek(0)

    first_url = str(reports[0].get("url") or "technical-seo")
    parsed = urlparse(first_url)
    host = (parsed.netloc or "technical-seo").replace(":", "-")
    path = (parsed.path or "").strip("/").replace("/", "_")
    url_label = host if not path else f"{host}_{path}"
    safe_label = "".join(
        char if char.isalnum() or char in "._-" else "_"
        for char in url_label
    ).strip("._-") or "technical-seo"
    timestamp = datetime.now().strftime("%Y%m%d_%H%M")

    return {
        "bytes": output.read(),
        "filename": f"{safe_label}_{timestamp}.pptx",
        "warnings": warnings,
    }


def build_multi_ppt_report_from_default_template(reports: list[dict]):
    return build_multi_ppt_report_from_template(
        reports=reports,
        template_bytes=_load_server_template_bytes(),
    )
